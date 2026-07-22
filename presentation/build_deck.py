"""Generate the Edge Negotiator project-pitch deck as a .pptx.

Built for a 3-minute pitch (~30s/slide) and the marking rubric:
Problem (25), Literature (35), Methodology + Evaluation (20), Presentation (20).
Glanceable one-line bullets, natural language, no em dashes.
"""
from pptx import Presentation
from pptx.util import Pt, Inches

prs = Presentation()
# Standard PowerPoint 16:9 widescreen
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def bullets(title, items):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    t = slide.shapes.title
    t.text = title
    t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.4), Inches(12.1), Inches(1.1)
    body = slide.placeholders[1]
    body.left, body.top, body.width, body.height = Inches(0.6), Inches(1.7), Inches(12.1), Inches(5.3)
    tf = body.text_frame
    tf.word_wrap = True
    for i, text in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        for r in p.runs:
            r.font.size = Pt(20)
    return slide


# Slide 1 - Title (clean)
s = prs.slides.add_slide(prs.slide_layouts[0])
st = s.shapes.title
st.text = "The Edge Negotiator"
st.left, st.top, st.width, st.height = Inches(0.8), Inches(2.1), Inches(11.7), Inches(1.3)
subph = s.placeholders[1]
subph.left, subph.top, subph.width, subph.height = Inches(0.8), Inches(3.5), Inches(11.7), Inches(3.4)
sub = subph.text_frame
sub.text = "A Trust-Preserving Coordination Layer for Signalised Junctions: Characterising a Self-Referential Coupling in Compromised-Insider Emergencies"
for line in [
    "UCL MSc Systems Engineering for IoT. Project pitch.",
    "Student 25153651, Ebenezer Veeraraju. Supervisors: Dr Akin Delibasi (UCL) and Lee Stott (Microsoft).",
]:
    p = sub.add_paragraph()
    p.text = line
    for r in p.runs:
        r.font.size = Pt(15)

# Slide 2 - Problem (Rubric I)
bullets("The problem", [
    "Emergency-vehicle preemption lets one signed message seize a junction's green phase, a lever a compromised insider key can pull without ever touching the road.",
    "The question: in a compromised-insider emergency, what can be held accountable mechanically, and what remains a human or counsel judgment that an on-device reasoner cannot faithfully automate?",
    "Signals are safety-critical infrastructure: preemption control is deployable only if a refused, uncorroborated claim is provably distinct from a cleared, physically corroborated one.",
    "Originality: no published account measures where the attack lever and the witness coverage that would catch it are the SAME variable.",
])

# Slide 3 - Literature (Rubric II, 35 pts)
bullets("Background and the gap", [
    "The accountability mechanism is not new: Certificate-Transparency-style, quorum-anchored, cross-audited logs (RFC 6962, CONIKS, A2M, TrInc, PeerReview, Kusters) are credited, not claimed novel.",
    "The stealthy attack class is False Data Injection (Liu, Ning and Reiter, 2011; Teixeira and Sandberg); unobservability's topology-dependence is owned by Kosut (2011) and Hendrickx (2014).",
    "The gap: nobody states or measures the coupling where the preemption attack's lever, the signal phase, is also the variable that gates honest-witness coverage.",
    "The nearest competitor is Traffic-R1; the differentiator is the accountability role plus the measured coupling, not a claim that the SLM beats a rule or that coordination optimises traffic.",
    "Honest limit: the coupling is a conditional lemma under explicit hypotheses (honest keys, quorum anchor), not an unconditional guarantee.",
])

# Slide 4 - Methodology (Rubric III, method)
bullets("Methodology", [
    "A deterministic real-time gate refuses signed-but-uncorroborated emergency preemption and clears physically corroborated real emergencies, on the real Euston Road (A501) corridor.",
    "One frozen Phi-4-mini SLM per junction (Microsoft Foundry Local) proposes a phase from the raw waiting-vehicle counts; a deterministic MaxPressure shield validates or overrides every decision.",
    "Fault output is split into two firewalled channels: a mechanically-verifiable provenance evidence pack with no verdict or accusation, and a counsel-gated internal triage note that names a key, never a person.",
    "The self-referential coupling is stated and measured directly: executing the preemption attack changes the phase, which changes the coverage that would corroborate or refute it.",
    "One severe, pre-registered measurement locates the honest boundary of that coupling on the real corridor, not an unconditional guarantee.",
])

# Slide 5 - Evaluation + progress (Rubric III, evaluation)
bullets("Evaluation and progress so far", [
    "Measured characterisation: the free-deviation classes the gate does not stop (sub-margin piggyback inflation, keyless transient spoofing) versus the out-of-scope axes (colluding keys, quorum compromise, coverage below threshold).",
    "SLM evaluation: citation-faithful legal-reasoning correctness against an un-rigged rule-to-text baseline, plus a characterised disambiguation classifier and a frozen-versus-hardened robustness tradeoff.",
    "Rigour: pre-registered seeds, bootstrap confidence intervals, multiple-comparison correction; the coordination effect itself is reported as a structural zero, not a headline benefit.",
    "Already working: the deterministic gate clears and refuses correctly on the fixture corridor, and a live SLM junction runs decisions end to end with zero failures.",
    "All code, scenario and documentation sit in a reproducible public repository.",
])

# Slide 6 - Roadmap, confidence, risks
bullets("Roadmap, confidence, and risks", [
    "Next: the Euston Road (A501) net build via netconvert, the pre-registered severe measurement of the coupling's boundary, and the SLM Job A/B/C characterisation.",
    "Confident: the deterministic gate and audit log are built and tested; the coupling claim is stated as a conditional lemma so it cannot overreach the evidence.",
    "Coupling risk: the measured boundary may be narrower than hoped; the honest boundary itself, not a guarantee, is the reported finding either way.",
    "SLM risk: a pre-committed kill criterion can demote the SLM's claim without removing the SLM's presence in the architecture.",
    "Data risk: calibrate demand from Department for Transport counts using SUMO routeSampler once the real net exists.",
])

out = __file__.rsplit("\\", 1)[0].rsplit("/", 1)[0] + "/25153651_EbenezerVeeraraju_ADelibasi_COMP0234_SysIoT_Pitch_2324.pptx"
prs.save(out)
print("saved:", out, "| slides:", len(prs.slides._sldIdLst))
